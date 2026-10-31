from pathlib import Path
import re
import unicodedata

import pandas as pd
import fitz
from docx import Document


QUESTION_PREFIX_RE = re.compile(
    r"^\s*(?:Q(?:UESTION)?\s*\d+\s*[.)\-:]?\s*)(.*)$",
    re.IGNORECASE,
)
NUMBERED_RE = re.compile(r"^\s*(\d+)\s*[.)\-:]\s+(.+?)\s*$")
OPTION_RE = re.compile(r"^\s*(?:[-•*]|[A-Za-z][.)]|\d+[.)])\s*(.+?)\s*$")
QUESTION_CUE_RE = re.compile(
    r"^(?:quel(?:le|s|les)?|combien|où|ou|quand|comment|pourquoi|qui|lequel|"
    r"indiquez|précisez|precisez|décrivez|decrivez|veuillez|avez-vous|"
    r"êtes-vous|etes-vous|est-ce|dans quelle|dans quel|what|which|where|"
    r"when|how|why|who|please|do you|are you)\b",
    re.IGNORECASE,
)



def extract_text(path: Path) -> str:
    """Extract selectable text from PDF, DOCX or TXT. OCR is not included."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        with fitz.open(path) as doc:
            text = "\n".join(page.get_text("text") for page in doc)
        if not text.strip():
            raise ValueError(
                "Aucun texte extractible dans ce PDF. Il semble numérisé : "
                "effectuez un OCR local puis réessayez."
            )
        return text
    if ext == ".docx":
        d = Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError("Formats acceptés : PDF, DOCX ou TXT.")


def _clean_line(line: str) -> str:
    return re.sub(r"\s+", " ", line).strip()


def parse_questionnaire(text: str):
    """Conservatively parse questions and answer options from extracted text."""
    lines = [_clean_line(line) for line in text.splitlines() if line.strip()]
    items = []
    current = None

    def flush():
        nonlocal current
        if current is not None:
            current["choices"] = list(dict.fromkeys(c for c in current["choices"] if c))
            items.append(current)
            current = None

    for line in lines:
        qprefix = QUESTION_PREFIX_RE.match(line)
        numbered = NUMBERED_RE.match(line)
        is_question = False
        label = line

        if qprefix:
            is_question = True
            label = qprefix.group(1).strip()
        elif numbered and ("?" in line or current is None or QUESTION_CUE_RE.match(numbered.group(2))):
            is_question = True
            label = numbered.group(2).strip()
        elif "?" in line and not (current is not None and OPTION_RE.match(line)):
            is_question = True
            label = re.sub(r"^\s*\d+\s*[.)\-:]\s*", "", line).strip()

        if is_question and label:
            flush()
            current = {"label": label, "choices": [], "hint": ""}
            continue

        option = OPTION_RE.match(line)
        if option and current is not None:
            current["choices"].append(option.group(1).strip())
            continue

        if current is not None:
            current["hint"] = (current["hint"] + " " + line).strip()

    flush()
    return items


RESERVED_NAMES = {"type", "label", "start", "today", "end", "deviceid", "subscriberid"}
MAX_VARIABLE_LENGTH = 32


def make_name(label, idx, max_length=MAX_VARIABLE_LENGTH):
    """Create a stable, ASCII lower_snake_case name; this is a proposal, not business approval."""
    text = unicodedata.normalize("NFKD", str(label))
    text = "".join(char for char in text if not unicodedata.combining(char))
    base = re.sub(r"[^A-Za-z0-9_]+", "_", text).lower()
    base = re.sub(r"_+", "_", base).strip("_")
    if not base:
        base = f"question_{idx}"
    if not base[0].isalpha():
        base = f"q_{base}"
    if base in RESERVED_NAMES:
        base = f"{base}_field"
    return base[:max_length].rstrip("_") or f"question_{idx}"[:max_length]


def _unique_name(base, seen, max_length=MAX_VARIABLE_LENGTH):
    """Return a unique name while keeping suffixes inside the configured length limit."""
    candidate = base[:max_length]
    suffix = 2
    while candidate.casefold() in seen:
        tail = f"_{suffix}"
        candidate = f"{base[:max_length - len(tail)].rstrip('_')}{tail}"
        suffix += 1
    seen.add(candidate.casefold())
    return candidate


def to_xlsform(items, out: Path, title="Questionnaire converti", form_id="questionnaire_converti"):
    survey, choices = [], []
    seen_names, seen_lists = set(), set()
    dictionary, traceability, decisions, logic = [], [], [], []

    for i, item in enumerate(items, 1):
        label = str(item.get("label", "")).strip()
        if not label:
            decisions.append({"source_ref": f"Q{i}", "issue": "Libellé source vide", "decision_required": "Retrouver et vérifier la question dans le document original", "status": "BLOQUÉ"})
            continue

        proposed_name = make_name(label, i)
        name = _unique_name(proposed_name, seen_names)
        options = [str(option).strip() for option in item.get("choices", []) if str(option).strip()]
        if options:
            list_base = make_name(f"list_{name}", i, max_length=45)
            list_name = _unique_name(list_base, seen_lists, max_length=45)
            qtype = f"select_one {list_name}"
            option_names = set()
            for k, option_label in enumerate(options, 1):
                option_base = make_name(option_label, k)
                option_name = _unique_name(option_base, option_names)
                choices.append({"list_name": list_name, "name": option_name, "label": option_label})
        else:
            list_name = ""
            qtype = "text"

        survey.append({
            "type": qtype, "name": name, "label": label,
            "hint": item.get("hint", ""), "required": "", "relevant": "",
            "constraint": "", "constraint_message": "", "calculation": "",
            "appearance": "", "choice_filter": "",
        })
        dictionary.append({
            "source_ref": f"Q{i}", "variable_name": name, "label_original": label,
            "data_type": qtype, "choice_list": list_name,
            "modalities_codes": "; ".join(f"{c['name']}={c['label']}" for c in choices if c["list_name"] == list_name) if list_name else "",
            "display_condition": "Non extraite automatiquement", "unit": "À déterminer dans la source",
            "definition": "À compléter après revue du questionnaire", "validation_status": "À VALIDER",
        })
        traceability.append({"source_ref": f"Q{i}", "source_location": "Position détectée dans l'ordre d'extraction; page/section non repérées", "source_label": label, "survey_row": len(survey) + 1, "variable_name": name, "status": "À VALIDER"})
        if proposed_name != name:
            decisions.append({"source_ref": f"Q{i}", "issue": f"Nom proposé normalisé ou désambiguïsé : {proposed_name} → {name}", "decision_required": "Confirmer le nom avant déploiement; documenter tout changement ultérieur", "status": "À VALIDER"})
        if not options:
            decisions.append({"source_ref": f"Q{i}", "issue": "Aucune modalité structurée détectée", "decision_required": "Confirmer si la réponse est libre ou si des modalités ont été manquées", "status": "À VALIDER"})
        logic.append({"source_ref": f"Q{i}", "variable_name": name, "relevant": "", "required": "", "constraint": "", "calculation": "", "choice_filter": "", "status": "NON TESTÉ", "notes": "La logique métier ne doit pas être inventée par la conversion."})

    if not survey:
        raise ValueError("Aucune question exploitable détectée ; XLSForm non généré.")

    settings = [{
        "form_title": title,
        "form_id": re.sub(r"[^A-Za-z0-9_-]", "_", form_id)[:60] or "questionnaire_converti",
        "version": "1",
        "default_language": "French (fr)",
    }]
    qa = [
        {"control": "Feuilles XLSForm principales générées", "status": "PASS", "details": "survey, choices, settings générées"},
        {"control": "Nommage lower_snake_case et limite 32 caractères", "status": "PASS" if all(len(row["name"]) <= MAX_VARIABLE_LENGTH and re.fullmatch(r"[a-z][a-z0-9_]*", row["name"]) for row in survey) else "BLOQUÉ", "details": "Contrôle automatique des noms de variables survey"},
        {"control": "Fidélité complète au document source", "status": "NON TESTÉ", "details": "Comparaison humaine obligatoire; extraction heuristique"},
        {"control": "Sauts, contraintes et calculs métier", "status": "NON TESTÉ", "details": "Ne sont pas reconstruits automatiquement"},
        {"control": "Importation Kobo/ODK/SurveyCTO", "status": "NON TESTÉ", "details": "Aucun test d'importation effectué"},
        {"control": "Validation métier et pilote", "status": "À VALIDER", "details": "Nécessaires avant toute déclaration PROD"},
    ]
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        pd.DataFrame(survey).to_excel(writer, sheet_name="survey", index=False)
        pd.DataFrame(choices, columns=["list_name", "name", "label"]).to_excel(writer, sheet_name="choices", index=False)
        pd.DataFrame(settings).to_excel(writer, sheet_name="settings", index=False)
        pd.DataFrame(dictionary).to_excel(writer, sheet_name="variable_dictionary", index=False)
        pd.DataFrame(traceability).to_excel(writer, sheet_name="source_traceability", index=False)
        pd.DataFrame(decisions, columns=["source_ref", "issue", "decision_required", "status"]).to_excel(writer, sheet_name="decisions", index=False)
        pd.DataFrame(logic).to_excel(writer, sheet_name="logic_register", index=False)
        pd.DataFrame(qa).to_excel(writer, sheet_name="qa_report", index=False)
        pd.DataFrame(columns=["test_id", "scenario", "expected_result", "observed_result", "status", "tester", "date", "notes"]).to_excel(writer, sheet_name="pilot_tests", index=False)
    return {"questions": len(survey), "choices": len(choices), "output": str(out), "documentation_sheets": 6}


def validate_items(items):
    warnings = []
    if not items:
        warnings.append("Aucune question détectée automatiquement. Revue humaine obligatoire.")
    seen_names = set()
    for i, item in enumerate(items, 1):
        label = str(item.get("label", "")).strip()
        if len(label) < 3:
            warnings.append(f"Question {i}: libellé trop court.")
        key = make_name(label, i)
        if key in seen_names:
            warnings.append(f"Question {i}: nom de variable potentiellement dupliqué.")
        seen_names.add(key)
        if item.get("choices") and len(item["choices"]) < 2:
            warnings.append(f"Question {i}: une seule modalité détectée ; vérifiez les choix.")
    return {"valid": len(items) > 0, "warnings": warnings}


def validate_xlsform(path: Path):
    """Validate workbook structure and common XLSForm naming/list-reference rules.

    This is a local structural check only; it does not replace pyxform or a
    platform-side validation/import test.
    """
    errors = []
    warnings = []
    path = Path(path)

    if path.suffix.lower() != ".xlsx":
        return {
            "valid": False,
            "errors": ["Le fichier doit être un classeur .xlsx."],
            "warnings": [],
        }

    try:
        workbook = pd.ExcelFile(path, engine="openpyxl")
        required_sheets = {"survey", "choices", "settings"}
        missing_sheets = sorted(required_sheets - set(workbook.sheet_names))
        if missing_sheets:
            return {
                "valid": False,
                "errors": [f"Feuilles obligatoires manquantes : {', '.join(missing_sheets)}."],
                "warnings": [],
            }

        survey = pd.read_excel(workbook, sheet_name="survey", dtype=str).fillna("")
        choices = pd.read_excel(workbook, sheet_name="choices", dtype=str).fillna("")
        settings = pd.read_excel(workbook, sheet_name="settings", dtype=str).fillna("")
    except (OSError, ValueError, ImportError, KeyError) as exc:
        return {
            "valid": False,
            "errors": [f"Impossible de lire le classeur XLSX : {exc}"],
            "warnings": [],
        }

    required_columns = {
        "survey": {"type", "name", "label"},
        "choices": {"list_name", "name", "label"},
        "settings": {"form_title", "form_id"},
    }
    frames = {"survey": survey, "choices": choices, "settings": settings}
    for sheet, required in required_columns.items():
        missing = sorted(required - set(frames[sheet].columns))
        if missing:
            errors.append(f"Colonnes manquantes dans '{sheet}' : {', '.join(missing)}.")
    if errors:
        return {"valid": False, "errors": errors, "warnings": warnings}

    if survey.empty:
        errors.append("La feuille 'survey' ne contient aucune question.")

    valid_identifier = re.compile(r"^[a-z][a-z0-9_]*$")
    valid_choice_code = re.compile(r"^[A-Za-z0-9_][A-Za-z0-9_-]*$")
    reserved_names = RESERVED_NAMES | {"instanceid", "username", "email", "phonenumber"}
    seen_survey_names = {}
    for excel_row, raw_name in enumerate(survey["name"], start=2):
        name = str(raw_name).strip()
        if not name:
            errors.append(f"survey.name est vide à la ligne Excel {excel_row}.")
            continue
        folded = name.casefold()
        if folded in seen_survey_names:
            errors.append(
                f"Nom de variable survey dupliqué (insensible à la casse) : "
                f"'{name}' aux lignes {seen_survey_names[folded]} et {excel_row}."
            )
        else:
            seen_survey_names[folded] = excel_row
        if not valid_identifier.fullmatch(name):
            errors.append(
                f"Nom survey invalide à la ligne {excel_row} : '{name}'. "
                "Utilisez lower_snake_case ASCII, commençant par une lettre."
            )
        if len(name) > MAX_VARIABLE_LENGTH:
            errors.append(
                f"Nom survey trop long à la ligne {excel_row} : '{name}' "
                f"({len(name)} caractères; maximum {MAX_VARIABLE_LENGTH})."
            )
        if folded in reserved_names:
            errors.append(f"Nom survey réservé à la ligne {excel_row} : '{name}'.")

    for column in ("label", "type"):
        for excel_row, value in enumerate(survey[column], start=2):
            if not str(value).strip():
                errors.append(f"Valeur vide dans survey.{column}, ligne Excel {excel_row}.")

    declared_lists = set()
    choice_pairs = set()
    if not choices.empty:
        for excel_row, row in choices.iterrows():
            list_name = str(row["list_name"]).strip()
            choice_name = str(row["name"]).strip()
            label = str(row["label"]).strip()
            actual_row = excel_row + 2
            if not list_name or not choice_name or not label:
                errors.append(
                    f"Valeur vide dans choices (ligne Excel {actual_row}); "
                    "list_name, name et label sont obligatoires."
                )
                continue
            declared_lists.add(list_name)
            pair = (list_name.casefold(), choice_name.casefold())
            if pair in choice_pairs:
                errors.append(
                    f"Modalité dupliquée dans choices à la ligne Excel {actual_row} : "
                    f"'{list_name}/{choice_name}'."
                )
            choice_pairs.add(pair)
            if not valid_identifier.fullmatch(list_name):
                errors.append(
                    f"Nom de liste invalide dans choices à la ligne {actual_row} : '{list_name}'."
                )
            if not valid_choice_code.fullmatch(choice_name):
                errors.append(
                    f"Code de modalité invalide dans choices à la ligne {actual_row} : '{choice_name}'."
                )

    used_lists = set()
    for excel_row, raw_type in enumerate(survey["type"], start=2):
        qtype = str(raw_type).strip()
        match = re.fullmatch(r"select_(one|multiple)\s+([A-Za-z][A-Za-z0-9_]*)", qtype)
        if qtype.startswith("select_") and not match:
            errors.append(
                f"Type de sélection mal formé dans survey.type ligne {excel_row} : '{qtype}'."
            )
        elif match:
            list_name = match.group(2)
            used_lists.add(list_name)
            if list_name not in declared_lists:
                errors.append(
                    f"survey.type ligne {excel_row} référence la liste absente '{list_name}'."
                )

    unused_lists = sorted(declared_lists - used_lists)
    if unused_lists:
        warnings.append(
            "Listes de choix déclarées mais non utilisées : " + ", ".join(unused_lists) + "."
        )

    if settings.empty:
        errors.append("La feuille 'settings' ne contient aucune ligne de configuration.")
    else:
        first = settings.iloc[0]
        for column in ("form_title", "form_id"):
            if not str(first[column]).strip():
                errors.append(f"settings.{column} est vide.")
        form_id = str(first["form_id"]).strip()
        if form_id and not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]*", form_id):
            errors.append(
                f"settings.form_id invalide : '{form_id}'. "
                "Il doit commencer par une lettre et contenir uniquement lettres, chiffres, '_' ou '-'."
            )
        if len(settings) > 1:
            warnings.append(
                "Plusieurs lignes détectées dans settings; seule la première est contrôlée."
            )

    return {"valid": not errors, "errors": errors, "warnings": warnings}
