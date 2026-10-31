from pathlib import Path

import fitz
import pandas as pd
import pytest
from docx import Document

from pdf_to_kobocollect import extract_text, parse_questionnaire, to_xlsform, validate_items, validate_xlsform


def test_numbered_answer_options_stay_with_their_question(tmp_path: Path):
    text = (
        "Q1. Quel est votre sexe ?\n1. Homme\n2. Femme\n"
        "Q2. Quel est votre âge ?\n1. Moins de 18 ans\n2. 18 ans ou plus\n"
    )
    items = parse_questionnaire(text)
    assert len(items) == 2
    assert items[0]["choices"] == ["Homme", "Femme"]
    assert items[1]["choices"] == ["Moins de 18 ans", "18 ans ou plus"]
    output = tmp_path / "form.xlsx"
    result = to_xlsform(items, output)
    assert result["questions"] == 2
    assert result["choices"] == 4
    survey = pd.read_excel(output, sheet_name="survey")
    choices = pd.read_excel(output, sheet_name="choices")
    assert len(survey) == 2
    assert len(choices) == 4
    assert survey["type"].str.startswith("select_one ").all()


def test_unmarked_question_with_question_mark_starts_new_item():
    items = parse_questionnaire(
        "Quel est votre pays ?\n- Djibouti\n- France\n"
        "Quelle est votre profession ?\n- Étudiant\n- Autre"
    )
    assert len(items) == 2
    assert items[0]["choices"] == ["Djibouti", "France"]
    assert items[1]["choices"] == ["Étudiant", "Autre"]


def test_duplicate_labels_get_unique_variable_and_choice_names(tmp_path: Path):
    items = parse_questionnaire("Q1. Choisissez une option ?\n1. Autre\n2. Autre")
    output = tmp_path / "duplicates.xlsx"
    to_xlsform(items, output)
    survey = pd.read_excel(output, sheet_name="survey")
    choices = pd.read_excel(output, sheet_name="choices")
    assert survey["name"].is_unique
    assert choices["name"].is_unique


def test_empty_questionnaire_is_not_exported(tmp_path: Path):
    with pytest.raises(ValueError, match="Aucune question"):
        to_xlsform([], tmp_path / "empty.xlsx")


def test_scanned_or_blank_pdf_returns_actionable_ocr_message(tmp_path: Path):
    pdf_path = tmp_path / "blank.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(pdf_path)
    with pytest.raises(ValueError, match="OCR local"):
        extract_text(pdf_path)


def test_validation_warns_when_only_one_choice_is_detected():
    result = validate_items([{"label": "Quel est votre statut ?", "choices": ["Étudiant"]}])
    assert result["valid"] is True
    assert any("une seule modalité" in warning for warning in result["warnings"])



def test_numbered_question_without_question_mark_is_detected():
    items = parse_questionnaire(
        "1. Quel est votre quartier\n- Balbala\n- Centre-ville\n"
        "2. Indiquez votre âge\n- Moins de 18 ans\n- 18 ans ou plus"
    )
    assert len(items) == 2
    assert items[0]["label"] == "Quel est votre quartier"
    assert items[0]["choices"] == ["Balbala", "Centre-ville"]
    assert items[1]["label"] == "Indiquez votre âge"
    assert items[1]["choices"] == ["Moins de 18 ans", "18 ans ou plus"]


def test_extract_text_from_txt_and_docx(tmp_path: Path):
    txt_path = tmp_path / "questionnaire.txt"
    txt_path.write_text("Quel est votre pays ?", encoding="utf-8")
    assert "Quel est votre pays ?" in extract_text(txt_path)

    docx_path = tmp_path / "questionnaire.docx"
    document = Document()
    document.add_paragraph("Quel est votre pays ?")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Oui"
    table.cell(0, 1).text = "Non"
    document.save(docx_path)
    extracted = extract_text(docx_path)
    assert "Quel est votre pays ?" in extracted
    assert "Oui | Non" in extracted



def test_generated_xlsform_passes_structural_validation(tmp_path: Path):
    items = parse_questionnaire("Q1. Quel est votre pays ?\\n1. Djibouti\\n2. France")
    output = tmp_path / "valid.xlsx"
    to_xlsform(items, output)
    report = validate_xlsform(output)
    assert report["valid"] is True
    assert report["errors"] == []


def test_xlsform_validator_detects_missing_choice_list(tmp_path: Path):
    output = tmp_path / "invalid.xlsx"
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([{"type": "select_one missing_list", "name": "country", "label": "Pays"}]).to_excel(
            writer, sheet_name="survey", index=False
        )
        pd.DataFrame([{"list_name": "other_list", "name": "yes", "label": "Oui"}]).to_excel(
            writer, sheet_name="choices", index=False
        )
        pd.DataFrame([{"form_title": "Test", "form_id": "test"}]).to_excel(
            writer, sheet_name="settings", index=False
        )
    report = validate_xlsform(output)
    assert report["valid"] is False
    assert any("missing_list" in error for error in report["errors"])


def test_naming_standard_transliterates_accents_and_limits_length():
    from pdf_to_kobocollect.engine import make_name
    assert make_name("État civil", 1) == "etat_civil"
    assert make_name("123 âge", 2).startswith("q_")
    assert len(make_name("Nombre de chameaux possédés actuellement dans le ménage", 3)) <= 32


def test_xlsform_includes_dictionary_traceability_and_qa_sheets(tmp_path: Path):
    output = tmp_path / "documented.xlsx"
    to_xlsform([{"label": "État civil ?", "choices": ["Marié", "Célibataire"]}], output)
    workbook = pd.ExcelFile(output)
    assert {"survey", "choices", "settings", "variable_dictionary", "source_traceability", "decisions", "logic_register", "qa_report", "pilot_tests"}.issubset(workbook.sheet_names)
    dictionary = pd.read_excel(output, sheet_name="variable_dictionary")
    assert dictionary.loc[0, "variable_name"] == "etat_civil"
    assert dictionary.loc[0, "validation_status"] == "À VALIDER"


def _write_minimal_xlsform(path, survey_rows, choice_rows=None, form_id="test_form"):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(survey_rows).to_excel(writer, sheet_name="survey", index=False)
        pd.DataFrame(
            choice_rows or [],
            columns=["list_name", "name", "label"],
        ).to_excel(writer, sheet_name="choices", index=False)
        pd.DataFrame([{"form_title": "Test", "form_id": form_id}]).to_excel(
            writer, sheet_name="settings", index=False
        )


def test_validator_rejects_invalid_and_overlong_variable_names(tmp_path):
    output = tmp_path / "invalid_names.xlsx"
    _write_minimal_xlsform(
        output,
        [
            {"type": "text", "name": "Age ménage", "label": "Âge"},
            {"type": "text", "name": "x" * 33, "label": "Autre"},
        ],
    )
    report = validate_xlsform(output)
    assert report["valid"] is False
    assert any("Nom survey invalide" in error for error in report["errors"])
    assert any("trop long" in error for error in report["errors"])


def test_validator_detects_case_insensitive_duplicate_names(tmp_path):
    output = tmp_path / "duplicate_case.xlsx"
    _write_minimal_xlsform(
        output,
        [
            {"type": "text", "name": "age_menage", "label": "Âge"},
            {"type": "integer", "name": "AGE_MENAGE", "label": "Âge du chef"},
        ],
    )
    report = validate_xlsform(output)
    assert any("dupliqué (insensible à la casse)" in error for error in report["errors"])


def test_validator_rejects_malformed_select_type_and_invalid_form_id(tmp_path):
    output = tmp_path / "bad_select.xlsx"
    _write_minimal_xlsform(
        output,
        [{"type": "select_one", "name": "pays", "label": "Pays"}],
        form_id="123 bad id",
    )
    report = validate_xlsform(output)
    assert report["valid"] is False
    assert any("mal formé" in error for error in report["errors"])
    assert any("form_id invalide" in error for error in report["errors"])


def test_validator_warns_about_unused_choice_lists(tmp_path):
    output = tmp_path / "unused_list.xlsx"
    _write_minimal_xlsform(
        output,
        [{"type": "text", "name": "commentaire", "label": "Commentaire"}],
        [{"list_name": "liste_oui_non", "name": "oui", "label": "Oui"}],
    )
    report = validate_xlsform(output)
    assert report["valid"] is True
    assert any("non utilisées" in warning for warning in report["warnings"])


def test_validator_rejects_non_xlsx_input(tmp_path):
    report = validate_xlsform(tmp_path / "questionnaire.csv")
    assert report["valid"] is False
    assert any(".xlsx" in error for error in report["errors"])


def test_validator_accepts_numeric_choice_codes(tmp_path):
    output = tmp_path / "numeric_choices.xlsx"
    _write_minimal_xlsform(
        output,
        [{"type": "select_one statut", "name": "statut", "label": "Statut"}],
        [
            {"list_name": "statut", "name": "1", "label": "Oui"},
            {"list_name": "statut", "name": "2", "label": "Non"},
        ],
    )
    report = validate_xlsform(output)
    assert report["valid"] is True
    assert report["errors"] == []
